# coding: utf-8
import hashlib,json,os,sys,tempfile,time,unittest
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
import project_pipeline as pp
class PipelineTests(unittest.TestCase):
 def setUp(self):
  (ROOT/'tmp/pipeline-tests').mkdir(parents=True,exist_ok=True)
  self.tmp=tempfile.TemporaryDirectory(dir=str(ROOT/'tmp/pipeline-tests'));self.root=Path(self.tmp.name);self.lib=self.root/'library';self.inbox=self.root/'inbox';self.archive=self.root/'archive';self.catalog=self.root/'catalog.json'
  pp.save(self.catalog,dict(projects=[dict(id='demo-01',name='测试项目',year=2025,award='待核对',unit='测试单位')]))
  self.lib.mkdir();pp.bootstrap(self.lib,self.inbox,self.catalog,self.archive);self.p=self.lib/'demo-01'
 def tearDown(self):self.tmp.cleanup()
 def source(self,name,text):
  p=self.inbox/'demo-01'/name;p.write_text(text,encoding='utf-8');return p
 def test_lossless_dedup_and_changes(self):
  path=self.source('研究.md','已有指标：12 mm。\n边界尚待核验。');first=pp.ingest(self.p,path);s=pp.state(self.p);meta=s['sources'][first['source_id']]
  self.assertEqual((self.p/meta['original']).read_bytes(),path.read_bytes());self.assertEqual(meta['sha256'],hashlib.sha256(path.read_bytes()).hexdigest())
  self.assertEqual(pp.ingest(self.p,path)['status'],'duplicate');self.assertEqual(pp.state(self.p)['revision'],s['revision'])
  path.write_text('修订后的记录：13 mm。');second=pp.ingest(self.p,path);self.assertNotEqual(first['source_id'],second['source_id']);self.assertEqual(len(pp.state(self.p)['sources']),2)
 def test_pinned_source_set_and_pagination(self):
  first=pp.ingest(self.p,self.source('a.txt','A'*4000));rev=first['revision'];pp.ingest(self.p,self.source('b.txt','新增资料'))
  result=pp.read_project(self.lib,'demo-01',level='sources',revision=rev);self.assertEqual(len(result['data']),1)
  result=pp.read_project(self.lib,'demo-01',level='source',source_id=first['source_id'],revision=rev,limit=1);self.assertEqual(len(result['data']),1);self.assertEqual(result['next_offset'],1)
  with self.assertRaises(ValueError):pp.read_project(self.lib,'demo-01',level='document',document='archive',revision=rev)
 def test_manual_document_is_not_overwritten(self):
  file=self.p/'01-项目档案.md';file.write_text('人工正文');pp.ingest(self.p,self.source('new.txt','新来源'))
  self.assertEqual(file.read_text(),'人工正文');self.assertTrue((self.p/'待合并/01-项目档案.md').exists());self.assertIn('01-项目档案.md',pp.state(self.p)['document_merge_required'])
 def test_html_scripts_removed_not_executed(self):
  p=self.source('page.html','<html><script>DO_NOT_RUN()</script><h1>题名</h1><p>指标 18 mm</p></html>');sections,status=pp.convert(p)
  self.assertEqual(status,'text_extracted');self.assertIn('18 mm',sections[0][1]);self.assertNotIn('DO_NOT_RUN',sections[0][1])
 def test_failures_and_unsupported_are_visible(self):
  self.source('broken.pdf','not a PDF');self.source('good.txt','真实段落');self.source('image.png','fake image, never OCR claimed')
  result=pp.scan(self.lib,self.inbox,0);self.assertEqual(result['changed'],3)
  statuses={x['status']for x in pp.state(self.p)['sources'].values()};self.assertEqual(statuses,{'conversion_failed','text_extracted','unsupported_needs_conversion'})
 def test_incremental_scan_and_unassigned(self):
  self.source('a.txt','一个原文');(self.inbox/'未归属.txt').write_text('未分配')
  first=pp.scan(self.lib,self.inbox,0);self.assertEqual(first['changed'],1);self.assertEqual(first['unassigned'],1)
  before=pp.state(self.p)['revision'];self.assertEqual(pp.scan(self.lib,self.inbox,0)['changed'],0);self.assertEqual(pp.state(self.p)['revision'],before)
 def test_unstable_file_deferred(self):
  self.source('copy.txt','复制中');self.assertEqual(pp.scan(self.lib,self.inbox,100)['changed'],0)
 def test_empty_pdf_needs_ocr(self):
  from pypdf import PdfWriter
  f=self.inbox/'demo-01/scan.pdf';writer=PdfWriter();writer.add_blank_page(width=200,height=200)
  with f.open('wb')as out:writer.write(out)
  result=pp.ingest(self.p,f);self.assertEqual(result['status'],'needs_ocr');self.assertEqual(pp.state(self.p)['sources'][result['source_id']]['chunk_count'],0)
 def test_docx_extract(self):
  from docx import Document
  doc=Document();doc.add_paragraph('原始设计记录');t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='指标';t.cell(0,1).text='8 mm'
  f=self.inbox/'demo-01/report.docx';doc.save(str(f));sections,status=pp.convert(f);self.assertIn('8 mm',sections[0][1]);self.assertIn('原始设计记录',sections[0][1])
 def test_no_traversal_or_arbitrary_revision(self):
  for pid in ['../demo-01','/tmp','x/y']:
   with self.assertRaises(ValueError):pp.project(self.lib,pid)
  with self.assertRaises(ValueError):pp.read_project(self.lib,'demo-01',revision='../../x')
  with self.assertRaises(ValueError):pp.read_project(self.lib,'demo-01',limit=21)
  with self.assertRaises(ValueError):pp.safe(self.p,'../secret')
 def test_public_url_blocks_private_and_unsupported_schemes(self):
  for url in ['file:///etc/passwd','http://user:pass@example.com']:
   with self.assertRaises(ValueError):pp.public_url(url)
  with patch.object(pp.socket,'getaddrinfo',return_value=[(2,1,6,'',('127.0.0.1',80))]):
   with self.assertRaises(ValueError):pp.public_url('http://local.invalid')
 def test_retry_keeps_failed_conversion_immutable(self):
  f=self.source('fail.pdf','invalid');first=pp.ingest(self.p,f);meta=pp.state(self.p)['sources'][first['source_id']];old=(self.p/meta['chunks']).read_bytes()
  with patch.object(pp,'convert',return_value=([('第1页','经修复的提取器取得文字')],'text_extracted')):second=pp.ingest(self.p,f,retry=True)
  self.assertNotEqual(first['source_id'],second['source_id']);self.assertEqual((self.p/meta['chunks']).read_bytes(),old)
  self.assertEqual(pp.state(self.p)['sources'][second['source_id']]['supersedes'],first['source_id'])
 def test_legacy_field_views_preserve_quotes_without_inventing(self):
  f=self.source('legacy.json',json.dumps(dict(desc='原库介绍12 mm',features='原库特点，不代表已证实'),ensure_ascii=False))
  pp.ingest(self.p,f);self.assertTrue(pp.attach_legacy(self.p,f));s=pp.state(self.p)
  self.assertEqual(s['legacy_extracts']['desc']['value'],'原库介绍12 mm')
  self.assertEqual(s['legacy_extracts']['features']['review_status'],'unverified_archive_excerpt')
  self.assertIn('原库介绍12 mm',(self.p/'01-项目档案.md').read_text());self.assertFalse(pp.attach_legacy(self.p,f))
 def test_bootstrap_repeat_does_not_reset_docs(self):
  f=self.p/'01-项目档案.md';f.write_text('我的正文');r=pp.bootstrap(self.lib,self.inbox,self.catalog,self.archive);self.assertEqual(r['created'],0);self.assertEqual(f.read_text(),'我的正文')
if __name__=='__main__':unittest.main()
