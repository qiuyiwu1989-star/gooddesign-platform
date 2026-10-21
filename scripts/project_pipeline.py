# coding: utf-8
"""024 local, lossless source pipeline and bounded archive reads. No LLM or MCP transport."""
import argparse, contextlib, datetime, fcntl, hashlib, ipaddress, json, mimetypes, os, re, shutil, socket, sys, time
import urllib.request, urllib.parse
from html.parser import HTMLParser
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
LIB=ROOT/'04-知识库/项目库'
INBOX=ROOT/'01-原始资料/项目投递箱'
TEMPLATE=ROOT/'04-知识库/模板/项目文档标准模板-v1'
VERSION='source-cleaner-0.1.0'
MAX_BYTES=30*1024*1024
MODULES=['项目概况','需求与系统','设计思路','指标与验证','专利与论文','历程与故事','应用与价值','质量与维护']
DOCS={'archive':'01-项目档案.md','library':'02-资料与解读.md','assessment':'03-档案评估.md'}
def now():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def sha(data):return hashlib.sha256(data).hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,v):
 p.parent.mkdir(parents=True,exist_ok=True);t=p.with_name(p.name+'.tmp')
 t.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');os.replace(t,p)
def valid_id(pid):
 if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,79}',pid):raise ValueError('项目ID无效')
 return pid
def safe(root,rel):
 p=(root/rel).resolve()
 if not p.is_relative_to(root.resolve()):raise ValueError('路径超出项目目录')
 return p
def project(lib,pid):
 p=safe(lib,valid_id(pid))
 if not (p/'数据/管道状态.json').is_file():raise ValueError('项目未建档，请先bootstrap')
 return p
def lock(lib):
 lib.mkdir(parents=True,exist_ok=True)
 @contextlib.contextmanager
 def locked():
  with (lib/'.pipeline.lock').open('a') as f:
   fcntl.flock(f,fcntl.LOCK_EX)
   try:yield
   finally:fcntl.flock(f,fcntl.LOCK_UN)
 return locked()
def state(p):return load(p/'数据/管道状态.json')
def log(lib,event,**values):
 with (lib/'管道日志.jsonl').open('a',encoding='utf-8') as f:f.write(json.dumps(dict(time=now(),event=event,**values),ensure_ascii=False)+'\n')
class PageText(HTMLParser):
 def __init__(self):super().__init__();self.skip=0;self.parts=[]
 def handle_starttag(self,tag,attrs):
  if tag in ['script','style','noscript']:self.skip+=1
  if not self.skip and tag in ['p','div','br','li','h1','h2','h3','tr','article']:self.parts.append('\n')
 def handle_endtag(self,tag):
  if tag in ['script','style','noscript']:self.skip=max(0,self.skip-1)
  if not self.skip and tag in ['p','div','li','h1','h2','h3','tr']:self.parts.append('\n')
 def handle_data(self,data):
  if not self.skip:self.parts.append(data)
def decode(data):
 for encoding in (['utf-16']if data.startswith((b'\xff\xfe',b'\xfe\xff'))else ['utf-8-sig','gb18030']):
  try:return data.decode(encoding)
  except UnicodeDecodeError:pass
 raise ValueError('无法可靠识别字符编码，待人工处理')
def convert(path):
 ext=path.suffix.lower();data=path.read_bytes()
 if ext in ['.txt','.md','.csv']:return [('正文',decode(data))], 'text_extracted'
 if ext=='.json':return [('JSON正文',json.dumps(json.loads(decode(data)),ensure_ascii=False,indent=2))], 'text_extracted'
 if ext in ['.html','.htm']:
  parser=PageText();parser.feed(decode(data));return [('网页正文',''.join(parser.parts))], 'text_extracted'
 if ext=='.pdf':
  from pypdf import PdfReader
  pages=[('第%d页'%(i+1),page.extract_text()or '')for i,page in enumerate(PdfReader(str(path)).pages)]
  empty=sum(not text.strip()for _,text in pages)
  return pages, 'needs_ocr'if empty==len(pages)else 'partial_needs_ocr'if empty else 'text_extracted'
 if ext=='.docx':
  from docx import Document
  doc=Document(str(path));rows=[p.text for p in doc.paragraphs]
  rows += [' | '.join(c.text for c in row.cells)for table in doc.tables for row in table.rows]
  return [('段落与表格（按类型汇集，非原始版式）','\n'.join(rows))], 'text_extracted'
 return [],'unsupported_needs_conversion'
def chunks(sections):
 out=[]
 for locator,text in sections:
  # Whitespace normalization only; no summarization, inferred facts, or instruction execution.
  lines=[re.sub(r'[ \t]+',' ',line).strip()for line in text.replace('\r\n','\n').splitlines()]
  clean='\n'.join(line for line in lines if line)
  for i in range(0,len(clean),1200):out.append(dict(id='B%04d'%(len(out)+1),locator=locator+' · 清洗文本字符%d–%d'%(i+1,min(i+1200,len(clean))),text=clean[i:i+1200]))
 return out

def render(p,s):
 from project_ownership import select
 meta=s['project'];sources=list(select(p,s['sources']).values());count=len(sources)
 heading='# '+meta['name']+'\n\n项目ID：'+meta['id']+'。原始材料按来源归档；尚未完成事实抽取和专业审核。\n\n'
 archive=heading+'## 档案总览\n\n现有来源 %d 份；这表示资料导入规模，不是研究完整度。\n\n'%count
 if meta['id']=='2025-00':archive+='已完成深研的嫦娥档案保留在`04-知识库/标杆案例/2025-00-嫦娥钻取采样/档案.html`。此目录为统一入库入口，不替代标杆。\n\n'
 archive+='## 旧库登记（待核验）\n\n年份：%s；奖项：%s；登记单位：%s。原始题录见来源目录。\n\n'%(meta.get('year',''),meta.get('award',''),meta.get('unit',''))
 for i,m in enumerate(MODULES,1):
  archive+='## %d %s\n\n'%(i,m)
  key={1:'desc',3:'features'}.get(i);quote=s.get('legacy_extracts',{}).get(key)
  if quote:
   archive+='旧库原字段摘录（待核验，未改写为平台结论）：\n\n'+''.join('> '+line+'\n'for line in quote['value'].splitlines())+'\n来源：['+quote['source_id']+']('+s['sources'][quote['source_id']]['document']+') · JSON字段 `'+key+'`。\n\n'
  else:archive+='待研究整合。资料先见[资料与解读](02-资料与解读.md)，不将原文自动改写为已核实事实。\n\n'

 library=heading+'## 来源目录\n\n以下为格式转换文档，尚不是论文或技术解读。原件、核读范围和清洗块可以逐级查看。\n\n'
 for x in sources:
  library+='- [%s · %s](%s)：%s；%d 块；待研究。\n'%(x['id'],x['title'].replace('\n',' '),x['document'],x['status'],x['chunk_count'])
 assessment=heading+'## 质量与审核\n\n质量分：未评。项目方/专家审核：未完成。来源转换不触发评分。\n\n## 模块整理情况\n\n'+''.join('- '+m+'：待研究。\n'for m in MODULES)+'\n## 待补事项\n\n- 按模块抽取事实并核对原始位置。\n- 形成设计机制、取舍与边界的有据解读。\n- 填写质量依据并安排审核。\n'
 conflicts=[]
 for key,text in [('archive',archive),('library',library),('assessment',assessment)]:
  dest=p/DOCS[key];existing=sha(dest.read_bytes())if dest.exists()else None;expected=s['generated'].get(key)
  if existing is not None and expected!=existing:
   target=p/'待合并'/DOCS[key];target.parent.mkdir(exist_ok=True);target.write_text(text,encoding='utf-8');conflicts.append(DOCS[key]);continue
  dest.write_text(text,encoding='utf-8');s['generated'][key]=sha(dest.read_bytes())
 s['document_merge_required']=conflicts
 return conflicts

def commit(p,s,reason):
 previous=state(p)
 from project_ownership import registry
 s['ownership_registry']=registry(p)
 s['ownership_revision']=s['ownership_registry'].get('version',0)
 render(p,s);s['revision']+=1;s['updated_at']=now()
 revision='ingest-%06d'%s['revision'];snap=p/'版本'/revision
 if snap.exists():raise ValueError('版本冲突，禁止覆盖')
 snap.mkdir(parents=True)
 save(snap/'snapshot.json',s)
 save(snap/'changes.json',dict(added_sources=sorted(set(s['sources'])-set(previous['sources'])),removed_sources=sorted(set(previous['sources'])-set(s['sources'])),changed_sources=[sid for sid in s['sources']if sid in previous['sources']and s['sources'][sid]!=previous['sources'][sid]],document_merge_required=s.get('document_merge_required',[])))
 save(snap/'manifest.json',dict(changes_sha256=sha((snap/'changes.json').read_bytes()),recorded_by=VERSION,revision=revision,recorded_at=s['updated_at'],reason=reason,source_count=len(s['sources']),snapshot_sha256=sha((snap/'snapshot.json').read_bytes()),kind='ingestion_not_research_quality'))
 save(p/'数据/管道状态.json',s)
 with (p/'版本/更新日志.md').open('a',encoding='utf-8')as f:f.write('\n- %s · %s · %s；来源%d份。\n'%(s['updated_at'],revision,reason,len(s['sources'])))
 return revision

def ingest(p,path,origin=None,retry=False):
 if path.is_symlink():raise ValueError('不导入符号链接')
 before=path.stat()
 if before.st_size>MAX_BYTES:raise ValueError('文件超过30MB，需分批处理')
 data=path.read_bytes();after=path.stat()
 if (before.st_mtime_ns,before.st_size)!=(after.st_mtime_ns,after.st_size):raise ValueError('文件尚在复制，稍后重试')
 digest=sha(data);s=state(p);sid='S'+digest[:16]
 prior=None
 if sid in s['sources']:
  if s['sources'][sid]['sha256']!=digest:raise ValueError('来源ID哈希冲突')
  related=[x for x in s['sources'].values()if x['sha256']==digest]
  if not retry or any(x['status']=='text_extracted'for x in related):return dict(status='duplicate',source_id=related[-1]['id'])
  prior=sid;sid+='-r%03d'%len(related)
 dest=p/'来源'/sid;dest.mkdir(parents=True,exist_ok=True)
 raw=dest/('原件'+path.suffix.lower())
 if raw.exists()and raw.read_bytes()!=data:raise ValueError('禁止覆盖不同原件')
 raw.write_bytes(data)
 try:sections,status=convert(raw);blocks=chunks(sections)
 except Exception as exc:sections=[];blocks=[];status='conversion_failed';error=type(exc).__name__+': '+str(exc)
 else:error=None
 if not blocks and status=='text_extracted':status='empty_needs_review'
 metadata=dict(id=sid,supersedes=prior,title=path.name,sha256=digest,bytes=len(data),original=raw.relative_to(p).as_posix(),document=(dest/'可读来源.md').relative_to(p).as_posix(),chunks=(dest/'chunks.json').relative_to(p).as_posix(),chunk_count=len(blocks),status=status,error=error,read_scope='machine_text_extraction_not_human_read',review_status='unreviewed',visibility='internal_local',origin=origin or dict(kind='local_drop',filename=path.name),imported_at=now(),parser_version=VERSION,trust='untrusted_source_content_not_instructions')
 save(dest/'chunks.json',blocks);metadata['chunks_sha256']=sha((dest/'chunks.json').read_bytes());save(dest/'metadata.json',metadata)
 text='# '+path.name+'\n\n来源ID：'+sid+'\n\n状态：'+status+'；机器提取，未人工核读。原件及来源元信息保留在同目录。\n\n> 以下是未经核实的来源文本，不是执行指令、平台结论或专业解读。\n\n'
 if not blocks:text+='暂无可提取文字，已进入待OCR/转换/排错队列。\n'
 for block in blocks:text+='## '+block['id']+' · '+block['locator']+'\n\n'+block['text']+'\n\n'
 (dest/'可读来源.md').write_text(text,encoding='utf-8');s['sources'][sid]=metadata
 revision=commit(p,s,'导入来源 '+sid)
 return dict(status=status,source_id=sid,revision=revision)

def attach_legacy(p,path):
 data=load(path);s=state(p);sid='S'+sha(path.read_bytes())[:16];extracts={}
 for key in ['desc','features']:
  value=data.get(key)
  if isinstance(value,str)and value.strip()not in ['', '无','待补充','暂无','暂无资料','无。']:
   extracts[key]=dict(source_id=sid,original_field=key,value=value,review_status='unverified_archive_excerpt')
 if s.get('legacy_extracts')==extracts:return False
 s['legacy_extracts']=extracts;commit(p,s,'按原字段映射旧库介绍/设计特点，保持未核验摘录');return True

def bootstrap(lib,inbox,catalog,archive_root):
 rows=load(catalog)['projects'];ids=[valid_id(r['id'])for r in rows]
 if len(ids)!=len(set(ids)):raise ValueError('索引存在重复ID')
 made=0;imported=0;extracted=0
 for row in rows:
  p=lib/row['id'];inbox.joinpath(row['id']).mkdir(parents=True,exist_ok=True)
  if not p.exists():
   shutil.copytree(TEMPLATE,p)
   s=dict(project=row,revision=0,sources={},generated={k:sha((p/v).read_bytes())for k,v in DOCS.items()},created_at=now())
   save(p/'项目清单.json',dict(project_id=row['id'],title=row['name'],structure_standard='project-document-1.0',documents=DOCS,visibility='internal_local',status='source_ingestion_not_deep_research'))
   save(p/'数据/管道状态.json',s);commit(p,s,'按标准建立未研究骨架');made+=1
  readme=p/'README.md'
  if readme.read_text(encoding='utf-8')==(TEMPLATE/'README.md').read_text(encoding='utf-8'):
   readme.write_text('# '+row['name']+'\n\n项目ID：'+row['id']+'。按统一标准建立的资料入库档案，旧库记录待核验，深度研究尚待开展。\n\n- [项目档案](01-项目档案.md)\n- [资料与解读](02-资料与解读.md)\n- [档案评估](03-档案评估.md)\n\n投递素材到平台01-原始资料/项目投递箱/'+row['id']+'/。管道生成来源文档与导入修订；人工改过的正文会保留，新稿进入待合并。\n',encoding='utf-8')
  for folder in ['数据','版本']:
   f=p/folder/'README.md'
   if f.read_text(encoding='utf-8')==(TEMPLATE/folder/'README.md').read_text(encoding='utf-8'):
    f.write_text('# '+folder+'\n\n本目录已接入project_pipeline.py的资料入库流程；运行状态是导入记录，不等于233字段深研完成或专业质量分。ingest修订保存来源集与完整性，研究正文与专业评分仍待后续整合。\n',encoding='utf-8')
  original=archive_root/'cases/detail'/ (row['id']+'.json')
  if original.is_file():
   result=ingest(p,original,dict(kind='legacy_archive',path=str(original.relative_to(ROOT))if original.is_relative_to(ROOT)else str(original),review_status='unverified_archive'))
   imported+=result['status']!='duplicate'
   extracted+=attach_legacy(p,original)
 save(lib/'项目索引.json',dict(version=VERSION,projects=[dict(id=r['id'],name=r['name'],path=r['id'],review_status='unreviewed',year=r.get('year'),award=r.get('award'),unit=r.get('unit'))for r in rows]))
 (inbox/'_待分配').mkdir(exist_ok=True)
 return dict(projects=len(rows),created=made,legacy_sources_imported=imported,legacy_field_views_updated=extracted)

def public_url(url):
 parsed=urllib.parse.urlsplit(url)
 if parsed.scheme not in ['http','https']or not parsed.hostname or parsed.username or parsed.password:raise ValueError('只支持无凭据的公开HTTP(S)来源')
 for result in socket.getaddrinfo(parsed.hostname,parsed.port or (443 if parsed.scheme=='https'else 80),type=socket.SOCK_STREAM):
  if not ipaddress.ip_address(result[4][0]).is_global:raise ValueError('抓取器不访问本机或内网地址')
 return url
class Redirect(urllib.request.HTTPRedirectHandler):
 def redirect_request(self,req,fp,code,msg,headers,newurl):
  public_url(newurl);return super().redirect_request(req,fp,code,msg,headers,newurl)
def fetch(lib,inbox,pid,url):
 p=project(lib,pid);public_url(url);opener=urllib.request.build_opener(Redirect())
 with opener.open(urllib.request.Request(url,headers={'User-Agent':'GoodDesignSourceArchive/0.1'}),timeout=20)as response:
  data=response.read(MAX_BYTES+1)
  if len(data)>MAX_BYTES:raise ValueError('响应超过30MB')
  content_type=response.headers.get_content_type();final=response.url
  ext={'text/html':'.html','application/pdf':'.pdf','application/json':'.json','text/plain':'.txt'}.get(content_type)
  if not ext:
   # Publishers often deliver PDF downloads as application/octet-stream. Preserve other formats too.
   if data.startswith(b'%PDF-'):ext='.pdf'
   else:
    candidate=Path(urllib.parse.urlsplit(final).path).suffix.lower()
    ext=candidate if candidate in ['.docx','.doc','.xlsx','.xls','.pptx','.ppt','.csv','.md','.txt','.png','.jpg','.jpeg','.webp','.xml','.zip']else '.bin'
  charset=response.headers.get_content_charset()
 target=inbox/pid/('web-'+sha(data)[:16]+ext);target.parent.mkdir(parents=True,exist_ok=True)
 target.write_bytes(data)
 save(target.with_name(target.name+'.origin.json'),dict(kind='web',requested_url=url,url=final,fetched_at=now(),content_type=content_type,declared_charset=charset))
 result=ingest(p,target,load(target.with_name(target.name+'.origin.json')))
 log(lib,'web_fetch',project_id=pid,url=url,final_url=final,**result)
 return result

def scan(lib,inbox,stable_seconds=2):
 results=[];unassigned=[];receipts=load(lib/'投递状态.json')if (lib/'投递状态.json').exists()else {}
 for path in sorted(inbox.rglob('*')):
  if not path.is_file()or path.is_symlink()or path.name.endswith('.origin.json')or path.name.startswith('.'):continue
  relative=path.relative_to(inbox);parts=relative.parts
  if len(parts)==1 and path.name=='README.md':continue
  if len(parts)<2 or parts[0]=='_待分配' or not (lib/parts[0]/'数据/管道状态.json').exists():unassigned.append(str(relative));continue
  stat=path.stat();signature=[stat.st_size,stat.st_mtime_ns];key=str(relative);receipt=receipts.get(key,{})
  if time.time()-stat.st_mtime<stable_seconds:continue
  if receipt.get('signature')==signature and (receipt.get('status')!='job_failed' or time.time()<receipt.get('retry_after',0)):continue
  try:
   p=project(lib,parts[0]);origin_path=path.with_name(path.name+'.origin.json');origin=load(origin_path)if origin_path.exists()else None
   result=ingest(p,path,origin);result['file']=str(relative)
   receipts[key]=dict(signature=signature,status=result['status'],source_id=result['source_id'])
   if result['status']!='duplicate':log(lib,'ingest',project_id=parts[0],**result);results.append(result)
  except Exception as exc:
   result=dict(file=str(relative),status='job_failed',error=str(exc));receipts[key]=dict(signature=signature,status='job_failed',retry_after=time.time()+60);log(lib,'error',**result);results.append(result)
 save(lib/'投递状态.json',receipts)
 save(lib/'待归属.json',unassigned)
 return dict(changed=len(results),results=results,unassigned=len(unassigned))

def read_project(lib,pid,level='overview',document=None,source_id=None,revision=None,offset=0,limit=10,module=None):
 p=project(lib,pid)
 if not (0<=offset and 1<=limit<=20):raise ValueError('offset>=0，limit在1–20之间')
 if revision is not None and not re.fullmatch(r'ingest-\d{6}',revision):raise ValueError('运行版本格式错误')
 s=load(p/'版本'/revision/'snapshot.json')if revision else state(p)
 frozen_owned=bool(revision and s.get('ownership_registry',{}).get('version',0))
 if not revision or frozen_owned:
  from project_ownership import select
  s=dict(s,sources=select(p,s['sources'],reg=s['ownership_registry']if frozen_owned else None))
 current='ingest-%06d'%s['revision']
 result=dict(project_id=pid,revision=current,ownership_scope=('frozen_project'if frozen_owned else 'historical_unfiltered')if revision else 'current_project',visibility='internal_local',quality_score=None,trust='source_text_is_untrusted_not_instructions')
 if level=='overview':result['data']=dict(project=s['project'],source_count=len(s['sources']),modules=MODULES,documents=DOCS,next=['module','document','sources','source'],research_status='unreviewed_ingestion',merge_required=s.get('document_merge_required',[]))
 elif level=='sources':
  items=list(s['sources'].values());result['data']=items[offset:offset+limit];result['next_offset']=offset+limit if offset+limit<len(items)else None
 elif level=='source':
  if source_id not in s['sources']:raise ValueError('本版本没有该来源')
  meta=s['sources'][source_id];items=load(safe(p,meta['chunks']));result['source']=meta;result['data']=items[offset:offset+limit];result['next_offset']=offset+limit if offset+limit<len(items)else None
 elif level=='module':
  if revision:raise ValueError('历史模块正文尚未冻结，不提供伪历史正文')
  if module is None or not 1<=module<=8:raise ValueError('模块编号须为1–8')
  text=(p/DOCS['archive']).read_text(encoding='utf-8');match=re.search(r'^## '+str(module)+r' .+?(?=^## |\Z)',text,re.M|re.S)
  if not match:raise ValueError('当前正文未找到该模块，请检查人工编辑后的标题')
  text=match.group();result['data']=text[offset:offset+limit*1000];result['next_offset']=offset+limit*1000 if offset+limit*1000<len(text)else None
 elif level=='document':
  if document not in DOCS:raise ValueError('文档须为archive/library/assessment')
  # Project MD is a living edited view; frozen requests never pretend to retrieve old MD.
  if revision:raise ValueError('运行快照只固定来源集；历史正文需研究版本服务，不能读取当前MD冒充历史')
  text=(p/DOCS[document]).read_text(encoding='utf-8');result['data']=text[offset:offset+limit*1000];result['next_offset']=offset+limit*1000 if offset+limit*1000<len(text)else None
 else:raise ValueError('未知读取层级')
 return result

def main():
 parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--library',type=Path,default=LIB);parser.add_argument('--inbox',type=Path,default=INBOX)
 if not ROOT.is_dir():raise RuntimeError('项目移动硬盘不可用，停止处理，不回退系统盘')
 sub=parser.add_subparsers(dest='command',required=True)
 b=sub.add_parser('bootstrap');b.add_argument('--catalog',type=Path,default=ROOT/'original-app/public/knowledge-data/project-index.json');b.add_argument('--archive',type=Path,default=ROOT/'99-历史项目快照/good-design/public')
 sub.add_parser('scan');w=sub.add_parser('watch');w.add_argument('--interval',type=float,default=4)
 retry_parser=sub.add_parser('retry');retry_parser.add_argument('project_id');retry_parser.add_argument('file',type=Path)
 search=sub.add_parser('search');search.add_argument('--query',default='');search.add_argument('--offset',type=int,default=0);search.add_argument('--limit',type=int,default=10)
 f=sub.add_parser('fetch');f.add_argument('project_id');f.add_argument('url')
 r=sub.add_parser('read');r.add_argument('project_id');r.add_argument('--level',choices=['overview','module','document','sources','source'],default='overview');r.add_argument('--document');r.add_argument('--module',type=int);r.add_argument('--source-id');r.add_argument('--revision');r.add_argument('--offset',type=int,default=0);r.add_argument('--limit',type=int,default=5)
 sub.add_parser('status');args=parser.parse_args();args.inbox.mkdir(parents=True,exist_ok=True)
 if args.command=='watch':
  if args.interval<1:parser.error('interval必须>=1')
  print('监视投递箱：'+str(args.inbox),flush=True)
  while True:
   if not ROOT.is_dir():raise RuntimeError('移动硬盘已离线，监听停止')
   with lock(args.library):result=scan(args.library,args.inbox)
   if result['changed']:print(json.dumps(result,ensure_ascii=False),flush=True)
   save(args.library/'监听状态.json',dict(pid=os.getpid(),heartbeat=now(),inbox=str(args.inbox)))
   time.sleep(args.interval)
 elif args.command=='read':print(json.dumps(read_project(args.library,args.project_id,args.level,args.document,args.source_id,args.revision,args.offset,args.limit,module=args.module),ensure_ascii=False))
 elif args.command=='search':
  if args.offset<0 or not 1<=args.limit<=20:parser.error('offset>=0，limit在1–20之间')
  rows=[r for r in load(args.library/'项目索引.json')['projects']if args.query.lower()in (r['id']+' '+r['name']).lower()]
  print(json.dumps(dict(data=rows[args.offset:args.offset+args.limit],total=len(rows),next_offset=args.offset+args.limit if args.offset+args.limit<len(rows)else None),ensure_ascii=False))
 elif args.command=='status':
  rows=load(args.library/'项目索引.json')['projects'];ss=[state(args.library/r['id'])for r in rows]
  print(json.dumps(dict(projects=len(ss),sources=sum(len(s['sources'])for s in ss),converted=sum(x['status']=='text_extracted'for s in ss for x in s['sources'].values()),needs_attention=sum(x['status']!='text_extracted'for s in ss for x in s['sources'].values()),merge_required=sum(bool(s.get('document_merge_required'))for s in ss),watcher=load(args.library/'监听状态.json')if (args.library/'监听状态.json').exists()else None),ensure_ascii=False))
 else:
  with lock(args.library):
   if args.command=='bootstrap':result=bootstrap(args.library,args.inbox,args.catalog,args.archive)
   elif args.command=='retry':result=ingest(project(args.library,args.project_id),args.file,retry=True)
   elif args.command=='fetch':result=fetch(args.library,args.inbox,args.project_id,args.url)
   else:result=scan(args.library,args.inbox)
  print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':
 try:main()
 except KeyboardInterrupt:sys.exit(0)
 except Exception as exc:
  print(json.dumps(dict(status='error',error=type(exc).__name__,message=str(exc)),ensure_ascii=False),file=sys.stderr);sys.exit(1)
